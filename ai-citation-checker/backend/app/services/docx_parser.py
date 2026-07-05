from __future__ import annotations
import io
import re
from dataclasses import dataclass, field
from docx import Document

REFERENCE_HEADINGS = {"references", "bibliography", "works cited", "reference list"}
STOP_HEADINGS = {"appendix", "appendices", "notes", "endnotes", "author note", "author notes"}

# A new reference entry starts with one of:
#   "Lastname, F."                (one or more authors with initials)
#   "Organisation Name (YYYY)"    also "Org. (YYYY)" and "Org. (n.d.)" —
#                                 APA puts a period after org names, and
#                                 undated web sources use (n.d.)
#   "lowercaseBrand. (YYYY|n.d.)" single-word brands like 'theCrag' — the
#                                 period is mandatory here so wrapped lines
#                                 ('the study (2019) …') never false-split
# Anything else is treated as a continuation of the previous entry.
_REF_START_RE = re.compile(
    r'^[A-ZÀ-Ɏ][\w\-\'‐‑\s&.]{0,80}?'
    r'(?:,\s*[A-ZÀ-Ɏ]\.|\.?\s*\((?:\d{4}|n\.d\.))'
    r'|^[a-zà-ɏ][\w\-]{0,30}\.\s*\((?:\d{4}|n\.d\.)'
)

# PDF copy artifact: a page/footnote number glued to the start of a reference
# ('12Wolter, B.') defeats the new-reference test above, so the paragraph is
# treated as a continuation and two references merge into one chimera entry.
# Only digits *directly* followed by an uppercase letter are junk — numbered
# list formats like '12. Wolter' or standalone numbers are not touched.
_LEADING_DIGIT_JUNK_RE = re.compile(r'^\d{1,4}(?=[A-ZÀ-Ɏ])')


@dataclass
class ReferenceParagraph:
    raw_text: str
    runs: list[tuple[str, bool]]      # (text, is_italic)
    has_hanging_indent: bool


@dataclass
class ParsedDocument:
    full_text: str
    body_text: str
    reference_paragraphs: list[ReferenceParagraph] = field(default_factory=list)


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"Failed to parse .docx: {exc}") from exc

    paragraphs = doc.paragraphs
    ref_start = _find_references_heading(paragraphs)

    body_paras = paragraphs[:ref_start] if ref_start is not None else paragraphs
    body_text = "\n".join(p.text for p in body_paras if p.text.strip())

    ref_paras: list[ReferenceParagraph] = []
    if ref_start is not None:
        for para in paragraphs[ref_start + 1:]:
            txt = para.text.strip()
            if not txt:
                continue
            if txt.lower() in STOP_HEADINGS or txt.lower().startswith("appendix"):
                break
            runs = [(r.text, _effective_italic(r, para)) for r in para.runs if r.text]
            para_text = para.text

            # Strip leading digit junk ('12Wolter, B.' → 'Wolter, B.') when
            # that's the only thing stopping the paragraph from being
            # recognised as a new reference.
            if not _REF_START_RE.match(txt):
                stripped = _LEADING_DIGIT_JUNK_RE.sub('', txt, count=1)
                if stripped != txt and _REF_START_RE.match(stripped):
                    txt = stripped
                    cleaned = _LEADING_DIGIT_JUNK_RE.sub('', para_text.lstrip(), count=1)
                    runs = _drop_leading_chars(runs, len(para_text) - len(cleaned))
                    para_text = cleaned

            fmt = para.paragraph_format
            style_fmt = para.style.paragraph_format if para.style else None
            direct_indent = fmt.first_line_indent
            style_indent = style_fmt.first_line_indent if style_fmt else None
            hanging = (
                (direct_indent is not None and direct_indent < 0)
                or (direct_indent is None and style_indent is not None and style_indent < 0)
            )

            # If this paragraph doesn't look like a new reference and we have a
            # previous one, treat it as a continuation (merge into previous).
            if ref_paras and not _REF_START_RE.match(txt):
                prev = ref_paras[-1]
                merged_text = prev.raw_text.rstrip() + " " + para_text.lstrip()
                ref_paras[-1] = ReferenceParagraph(
                    raw_text=merged_text,
                    runs=prev.runs + runs,
                    has_hanging_indent=prev.has_hanging_indent,
                )
                continue

            ref_paras.append(ReferenceParagraph(
                raw_text=para_text,
                runs=runs,
                has_hanging_indent=hanging,
            ))

    full_text = "\n".join(p.text for p in paragraphs if p.text.strip())
    return ParsedDocument(
        full_text=full_text,
        body_text=body_text,
        reference_paragraphs=ref_paras,
    )


def _drop_leading_chars(runs: list[tuple[str, bool]], n: int) -> list[tuple[str, bool]]:
    """Drop the first ``n`` characters from a run list, preserving italics —
    keeps runs aligned with raw_text after leading junk is stripped."""
    out: list[tuple[str, bool]] = []
    for text, italic in runs:
        if n >= len(text):
            n -= len(text)
            continue
        if n:
            text = text[n:]
            n = 0
        out.append((text, italic))
    return out


def _find_references_heading(paragraphs) -> int | None:
    for i, para in enumerate(paragraphs):
        if para.text.strip().lower() in REFERENCE_HEADINGS:
            return i
    return None


def _style_chain_italic(style) -> bool | None:
    """Walk a style and its ``base_style`` ancestors looking for an explicit
    italic setting. Returns True / False / None (not set anywhere in chain)."""
    seen: set[int] = set()
    cur = style
    while cur is not None and id(cur) not in seen:
        seen.add(id(cur))
        font = getattr(cur, "font", None)
        italic = getattr(font, "italic", None) if font is not None else None
        if italic is not None:
            return bool(italic)
        cur = getattr(cur, "base_style", None)
    return None


def _effective_italic(run, paragraph) -> bool:
    """Resolve a run's italic state through style inheritance.

    python-docx's ``run.italic`` returns:
      - True / False : italic was toggled directly on the run.
      - None         : not set on the run — inherit from character / paragraph style.

    Reading ``run.italic`` naively (== True) misses italics applied via a
    character style (e.g. ``Emphasis``) or via the paragraph style — including
    when those styles themselves inherit italic from a ``base_style``. We walk
    the cascade so all of these still count as italic.
    """
    if run.italic is not None:
        return bool(run.italic)
    val = _style_chain_italic(getattr(run, "style", None))
    if val is not None:
        return val
    val = _style_chain_italic(getattr(paragraph, "style", None))
    if val is not None:
        return val
    return False
